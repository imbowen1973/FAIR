"""Merged table cells, end to end, and the import cleaning they exposed.

A merge is written in place: a cell that is exactly "<" belongs to the
cell on its left, exactly "^" to the cell above. The round trip is the
test that matters: a deck with real PowerPoint merges, imported to
markdown and rendered again, must come back with the same merges.
"""

from __future__ import annotations

import zipfile
from pathlib import Path

import pytest
from lxml import etree

from edufair_renderer.ingest import _clean, ingest_pptx
from edufair_renderer.parser import SessionParseError
from edufair_renderer.render import render_session
from edufair_renderer.tablefit import fit_table

REPO = Path(__file__).resolve().parents[2]
EXAMPLES = REPO / "examples"
A_NS = "http://schemas.openxmlformats.org/drawingml/2006/main"


def _render(tmp_path, body: str):
    md = tmp_path / "s.md"
    md.write_text(
        "---\nsession: t1\ntitle: T\nversion: 0.1.0\n---\n\n"
        "--- slide\nid: s-01\nlayout: Full\ntitle: Merged\n"
        f"full:\n{body}---\n",
        encoding="utf-8",
    )
    out = render_session(
        session_path=md,
        template_path=EXAMPLES / "template.pptx",
        layout_map_path=EXAMPLES / "layout-map.yaml",
        out_dir=tmp_path / "out",
    )
    with zipfile.ZipFile(out["pptx"]) as z:
        return etree.fromstring(z.read("ppt/slides/slide1.xml"))


def _cells(root) -> list[list[etree._Element]]:
    return [list(tr.iter(f"{{{A_NS}}}tc")) for tr in root.iter(f"{{{A_NS}}}tr")]


def _text(tc) -> str:
    return "".join(t.text or "" for t in tc.iter(f"{{{A_NS}}}t"))


MERGED = (
    "  type: table\n"
    "  header: [Stage, Timing, '<']\n"
    "  rows:\n"
    "  - [Receive, Day 1, Within an hour]\n"
    "  - ['^', Day 2, Before dispatch]\n"
)


def test_markers_render_as_real_merges(tmp_path):
    cells = _cells(_render(tmp_path, MERGED))
    # "Timing" spans two columns; the cell it covers is hMerge'd and empty.
    assert cells[0][1].get("gridSpan") == "2"
    assert cells[0][2].get("hMerge") == "1"
    assert _text(cells[0][1]) == "Timing" and _text(cells[0][2]) == ""
    # "Receive" spans two rows; the cell below is vMerge'd and empty.
    assert cells[1][0].get("rowSpan") == "2"
    assert cells[2][0].get("vMerge") == "1"
    assert _text(cells[1][0]) == "Receive" and _text(cells[2][0]) == ""


def test_a_block_accepts_either_marker_inside(tmp_path):
    body = (
        "  type: table\n  rows:\n"
        "  - [Both, '<', c]\n"
        "  - ['^', '^', d]\n"   # interior written '^' ...
    )
    cells = _cells(_render(tmp_path, body))
    assert cells[0][0].get("gridSpan") == "2" and cells[0][0].get("rowSpan") == "2"


def test_escaped_markers_are_literal(tmp_path):
    body = "  type: table\n  rows:\n  - ['\\<', '\\^']\n  - [a, b]\n"
    cells = _cells(_render(tmp_path, body))
    assert [_text(c) for c in cells[0]] == ["<", "^"]
    assert all(c.get("gridSpan") is None for row in cells for c in row)


@pytest.mark.parametrize(
    "rows, message",
    [
        ("  - ['<', a]\n  - [b, c]\n", "nothing to its left"),
        ("  - ['^', a]\n  - [b, c]\n", "nothing above it"),
        ("  - [a, '<']\n  - ['^', b]\n", "not a rectangle"),
    ],
)
def test_malformed_merges_are_refused(tmp_path, rows, message):
    with pytest.raises(SessionParseError, match=message):
        _render(tmp_path, "  type: table\n  rows:\n" + rows)


def test_spanned_text_gets_the_combined_width():
    """A sentence across three columns must not shrink the table as if it
    had to fit in one of them."""
    sentence = "This guidance applies to every product in the cold chain"
    W, H = 8 * 914400, int(4.5 * 914400)
    merged = fit_table(
        None, [[sentence, "", ""], ["a", "b", "c"]], W, H, spans={(0, 0): (1, 3)}
    )
    assert merged.font_pt == 18.0


def test_import_round_trips_real_merges(tmp_path):
    from pptx import Presentation
    from pptx.util import Inches

    prs = Presentation()
    s = prs.slides.add_slide(prs.slide_layouts[6])
    t = s.shapes.add_textbox(Inches(0.8), Inches(0.4), Inches(8), Inches(0.9))
    t.text_frame.paragraphs[0].add_run().text = "Merged"
    table = s.shapes.add_table(3, 3, Inches(1), Inches(1.8), Inches(8), Inches(2)).table
    values = [["Stage", "Timing", ""], ["Receive", "Day 1", "Within an hour"],
              ["", "Day 2", "Before dispatch"]]
    for r, row in enumerate(values):
        for c, v in enumerate(row):
            table.cell(r, c).text = v
    table.cell(0, 1).merge(table.cell(0, 2))
    table.cell(1, 0).merge(table.cell(2, 0))
    deck = tmp_path / "merged.pptx"
    prs.save(str(deck))

    md = ingest_pptx(deck, "x").slides_md
    import yaml

    full = yaml.safe_load(md.split("--- slide", 1)[1].split("\n---", 1)[0])["full"]
    assert full["header"] == ["Stage", "Timing", "<"]
    assert full["rows"][1] == ["^", "Day 2", "Before dispatch"]

    path = tmp_path / "x.md"
    path.write_text(md, encoding="utf-8")
    out = render_session(
        session_path=path,
        template_path=EXAMPLES / "template.pptx",
        layout_map_path=EXAMPLES / "layout-map.yaml",
        out_dir=tmp_path / "out",
    )
    with zipfile.ZipFile(out["pptx"]) as z:
        cells = _cells(etree.fromstring(z.read("ppt/slides/slide1.xml")))
    assert cells[0][1].get("gridSpan") == "2"
    assert cells[1][0].get("rowSpan") == "2"
    assert _text(cells[1][0]) == "Receive"


def test_import_escapes_cells_that_are_literally_markers(tmp_path):
    from pptx import Presentation
    from pptx.util import Inches

    prs = Presentation()
    s = prs.slides.add_slide(prs.slide_layouts[6])
    table = s.shapes.add_table(2, 2, Inches(1), Inches(1.8), Inches(8), Inches(2)).table
    for (r, c), v in {(0, 0): "<", (0, 1): "^", (1, 0): "a", (1, 1): "b"}.items():
        table.cell(r, c).text = v
    deck = tmp_path / "lit.pptx"
    prs.save(str(deck))
    md = ingest_pptx(deck, "x").slides_md
    assert "\\<" in md and "\\^" in md


@pytest.mark.parametrize(
    "raw, cleaned",
    [
        ("oxygen", "oxygen"),           # was "xygen"
        ("one dose only", "one dose only"),
        ("-20 °C", "-20 °C"),           # was "20 °C": a sign lost
        ("—", "—"),                     # was "": an n/a mark lost
        ("o  bulleted item", "bulleted item"),
        ("• real bullet", "real bullet"),
        ("- dash bullet", "dash bullet"),
    ],
)
def test_bullet_stripping_needs_whitespace(raw, cleaned):
    assert _clean(raw) == cleaned


def test_table_cells_keep_their_leading_signs(tmp_path):
    from pptx import Presentation
    from pptx.util import Inches

    prs = Presentation()
    s = prs.slides.add_slide(prs.slide_layouts[6])
    table = s.shapes.add_table(2, 2, Inches(1), Inches(1.8), Inches(8), Inches(2)).table
    for (r, c), v in {(0, 0): "Store", (0, 1): "Range", (1, 0): "Freezer",
                      (1, 1): "- 25 to -15 °C"}.items():
        table.cell(r, c).text = v
    deck = tmp_path / "signs.pptx"
    prs.save(str(deck))
    assert "- 25 to -15 °C" in ingest_pptx(deck, "x").slides_md
