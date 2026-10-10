"""Pictures on import: found wherever they sit, bound to a placeholder.

A legacy deck puts pictures in placeholders, loose on the slide, as a
shape's fill, or behind everything as the background, and lays text
beside them, under them, or on top of them. Import has to find every
one, bind the slide's picture to a placeholder (Picture alone, Split
beside text on the side it was), say when text was overlaid, and file
the media under names that cannot collide with another import.
"""

from __future__ import annotations

import io
import zipfile
from pathlib import Path

import pytest
import yaml
from PIL import Image
from pptx import Presentation
from pptx.util import Inches

from edufair_renderer.assets import MAX_IMAGE_BYTES
from edufair_renderer.ingest import ingest_pptx, write_block
from edufair_renderer.render import render_session

REPO = Path(__file__).resolve().parents[2]
EXAMPLES = REPO / "examples"


def _png(path: Path, size=(1200, 800), colour=(30, 90, 140)) -> Path:
    Image.new("RGB", size, colour).save(path)
    return path


def _slides(md: str) -> list[dict]:
    return [yaml.safe_load(c.split("\n---", 1)[0]) for c in md.split("--- slide")[1:]]


def _deck(tmp_path, build) -> Path:
    prs = Presentation()
    build(prs, tmp_path)
    path = tmp_path / "deck.pptx"
    prs.save(str(path))
    return path


def _text(slide, text, left, top, width=4, height=1):
    box = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    box.text_frame.text = text
    return box


def test_picture_in_a_placeholder_is_found(tmp_path):
    def build(prs, d):
        s = prs.slides.add_slide(prs.slide_layouts[8])  # Picture with Caption
        s.shapes.title.text = "In a placeholder"
        ph = next(p for p in s.placeholders if p.placeholder_format.type == 18)
        ph.insert_picture(str(_png(d / "a.png")))

    result = ingest_pptx(_deck(tmp_path, build), "b")
    (slide,) = _slides(result.slides_md)
    assert slide["layout"] == "Picture"
    assert slide["picture"]["src"] in result.media


@pytest.mark.parametrize("picture_left", [True, False])
def test_picture_beside_text_keeps_its_side(tmp_path, picture_left):
    def build(prs, d):
        s = prs.slides.add_slide(prs.slide_layouts[6])
        _text(s, "Beside", 0.8, 0.4, 8)
        pic_x, text_x = (0.6, 5.2) if picture_left else (5.2, 0.6)
        s.shapes.add_picture(str(_png(d / "a.png")), Inches(pic_x), Inches(1.8), width=Inches(4))
        body = _text(s, "Check the seal", text_x, 1.8, 4, 3)
        body.text_frame.add_paragraph().text = "Record the reading"

    (slide,) = _slides(ingest_pptx(_deck(tmp_path, build), "b").slides_md)
    assert slide["layout"] == "Split"
    pic_side, text_side = ("left", "right") if picture_left else ("right", "left")
    assert slide[pic_side]["type"] == "image"
    assert slide[text_side]["items"] == ["Check the seal", "Record the reading"]


def test_short_line_under_a_picture_is_its_caption(tmp_path):
    def build(prs, d):
        s = prs.slides.add_slide(prs.slide_layouts[6])
        _text(s, "A week of readings", 0.8, 0.4, 8)
        s.shapes.add_picture(str(_png(d / "a.png")), Inches(1.5), Inches(1.6), width=Inches(6))
        _text(s, "The red line marks the threshold.", 1.5, 5.8, 6, 0.6)

    (slide,) = _slides(ingest_pptx(_deck(tmp_path, build), "b").slides_md)
    assert slide["layout"] == "Picture"
    assert slide["caption"] == {"type": "p", "text": "The red line marks the threshold."}


def test_a_list_beside_a_picture_is_never_a_caption(tmp_path):
    """The bug: a short list next to a picture became its caption."""
    def build(prs, d):
        s = prs.slides.add_slide(prs.slide_layouts[6])
        _text(s, "Receiving", 0.8, 0.4, 8)
        s.shapes.add_picture(str(_png(d / "a.png")), Inches(0.6), Inches(1.8), width=Inches(4))
        body = _text(s, "Seal", 5.2, 1.8, 4, 3)
        for line in ("Reading", "Quarantine"):
            body.text_frame.add_paragraph().text = line

    (slide,) = _slides(ingest_pptx(_deck(tmp_path, build), "b").slides_md)
    assert slide["layout"] == "Split"
    assert "caption" not in slide


def test_overlaid_text_becomes_the_caption_and_is_reported(tmp_path):
    def build(prs, d):
        s = prs.slides.add_slide(prs.slide_layouts[6])
        s.shapes.add_picture(str(_png(d / "a.png")), 0, 0,
                             width=prs.slide_width, height=prs.slide_height)
        _text(s, "Every degree matters", 1, 5.5, 8, 1)

    result = ingest_pptx(_deck(tmp_path, build), "b")
    (slide,) = _slides(result.slides_md)
    assert slide["layout"] == "Picture"
    assert slide["caption"]["text"] == "Every degree matters"
    notes = result.report[0]["notes"]
    assert any("was overlaid on the picture" in n and "Every degree matters" in n for n in notes)


def test_a_title_over_the_picture_is_reported(tmp_path):
    def build(prs, d):
        s = prs.slides.add_slide(prs.slide_layouts[5])  # Title Only
        s.shapes.title.text = "Over the photo"
        s.shapes.add_picture(str(_png(d / "a.png")), 0, 0,
                             width=prs.slide_width, height=prs.slide_height)

    result = ingest_pptx(_deck(tmp_path, build), "b")
    assert any("was set over the picture" in n for n in result.report[0]["notes"])


def test_a_picture_used_as_a_shape_fill_is_found(tmp_path):
    from pptx.enum.shapes import MSO_SHAPE

    def build(prs, d):
        s = prs.slides.add_slide(prs.slide_layouts[6])
        _text(s, "Filled", 0.8, 0.4, 8)
        shape = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(1), Inches(1.6),
                                   Inches(6), Inches(4))
        # python-pptx has no API for a picture fill; build the blipFill
        # the way PowerPoint writes one.
        image_part, rid = s.part.get_or_add_image_part(str(_png(d / "f.png")))
        sppr = shape._element.spPr
        for old in sppr.findall("{http://schemas.openxmlformats.org/drawingml/2006/main}solidFill"):
            sppr.remove(old)
        from lxml import etree
        a = "http://schemas.openxmlformats.org/drawingml/2006/main"
        r = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
        fill = etree.SubElement(sppr, f"{{{a}}}blipFill")
        etree.SubElement(fill, f"{{{a}}}blip").set(f"{{{r}}}embed", rid)
        etree.SubElement(etree.SubElement(fill, f"{{{a}}}stretch"), f"{{{a}}}fillRect")
        prst = sppr.find(f"{{{a}}}prstGeom")
        prst.addnext(fill)

    result = ingest_pptx(_deck(tmp_path, build), "b")
    (slide,) = _slides(result.slides_md)
    assert slide["layout"] == "Picture"
    assert slide["picture"]["src"] in result.media


def test_slide_background_is_saved_not_placed(tmp_path):
    def build(prs, d):
        s = prs.slides.add_slide(prs.slide_layouts[6])
        _text(s, "On a photographic background", 0.8, 0.4, 8)
        _text(s, "Body text only", 0.8, 2, 8, 2)
        s.background.fill.solid()  # make the <p:bg> exist, then swap in a picture
        from lxml import etree
        p = "http://schemas.openxmlformats.org/presentationml/2006/main"
        a = "http://schemas.openxmlformats.org/drawingml/2006/main"
        r = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
        _, rid = s.part.get_or_add_image_part(str(_png(d / "bg.png")))
        bgpr = s._element.find(f"{{{p}}}cSld/{{{p}}}bg/{{{p}}}bgPr")
        for child in list(bgpr):
            if child.tag == f"{{{a}}}solidFill":
                bgpr.remove(child)
        fill = etree.Element(f"{{{a}}}blipFill")
        etree.SubElement(fill, f"{{{a}}}blip").set(f"{{{r}}}embed", rid)
        bgpr.insert(0, fill)

    result = ingest_pptx(_deck(tmp_path, build), "b")
    (slide,) = _slides(result.slides_md)
    assert slide["layout"] == "Full"  # placed as text, not as a picture slide
    assert len(result.media) == 1
    assert any("background picture saved" in n for n in result.report[0]["notes"])


def test_a_logo_on_every_slide_is_branding_not_content(tmp_path):
    def build(prs, d):
        logo = _png(d / "logo.png", (200, 200), (200, 30, 30))
        for n in range(4):
            s = prs.slides.add_slide(prs.slide_layouts[6])
            _text(s, f"Slide {n}", 0.8, 0.4, 8)
            _text(s, "Some body text", 0.8, 2, 8, 2)
            s.shapes.add_picture(str(logo), Inches(8.8), Inches(6.6), width=Inches(0.8))

    result = ingest_pptx(_deck(tmp_path, build), "b")
    assert result.media == {}
    assert all(s["layout"] == "Full" for s in _slides(result.slides_md))
    assert any("dropped as the old deck's branding" in n for n in result.report[0]["notes"])


def test_media_names_cannot_collide_across_imports(tmp_path):
    def build_with(colour):
        def build(prs, d):
            s = prs.slides.add_slide(prs.slide_layouts[6])
            s.shapes.add_picture(str(_png(d / f"{colour}.png", colour=colour)), 0, 0,
                                 width=Inches(4))
        return build

    (tmp_path / "a").mkdir()
    (tmp_path / "b").mkdir()
    first = ingest_pptx(_deck(tmp_path / "a", build_with((200, 0, 0))), "blk")
    second = ingest_pptx(_deck(tmp_path / "b", build_with((0, 0, 200))), "blk")
    # Same block, same file name, same slide number: still different files.
    assert set(first.media).isdisjoint(second.media)
    (name,) = first.media
    assert name.startswith("media/deck-s01-")


def test_crop_is_applied_so_media_matches_the_slide(tmp_path):
    def build(prs, d):
        img = Image.new("RGB", (1000, 500), (220, 0, 0))
        img.paste((0, 0, 220), (500, 0, 1000, 500))  # right half blue
        img.save(d / "halves.png")
        s = prs.slides.add_slide(prs.slide_layouts[6])
        pic = s.shapes.add_picture(str(d / "halves.png"), 0, 0, width=Inches(4))
        pic.crop_left = 0.5  # the slide shows only the blue half

    result = ingest_pptx(_deck(tmp_path, build), "b")
    (data,) = result.media.values()
    im = Image.open(io.BytesIO(data)).convert("RGB")
    r, g, b = im.getpixel((im.width // 2, im.height // 2))
    assert b > 150 and r < 60


def test_large_pictures_fit_under_the_asset_cap(tmp_path):
    import os

    def build(prs, d):
        noise = Image.frombytes("RGB", (2600, 1800), os.urandom(2600 * 1800 * 3))
        noise.save(d / "noise.png")
        s = prs.slides.add_slide(prs.slide_layouts[6])
        s.shapes.add_picture(str(d / "noise.png"), 0, 0, width=Inches(6))

    result = ingest_pptx(_deck(tmp_path, build), "b")
    (data,) = result.media.values()
    assert len(data) <= MAX_IMAGE_BYTES


def test_imported_pictures_render(tmp_path):
    """Round trip: import, write the block, render it -- the media is
    where the slides say it is, and the pictures are in the deck."""
    def build(prs, d):
        s = prs.slides.add_slide(prs.slide_layouts[8])
        s.shapes.title.text = "In a placeholder"
        ph = next(p for p in s.placeholders if p.placeholder_format.type == 18)
        ph.insert_picture(str(_png(d / "a.png")))
        s = prs.slides.add_slide(prs.slide_layouts[6])
        _text(s, "Beside", 0.8, 0.4, 8)
        s.shapes.add_picture(str(_png(d / "b.png", colour=(0, 120, 0))),
                             Inches(0.6), Inches(1.8), width=Inches(4))
        body = _text(s, "One", 5.2, 1.8, 4, 3)
        body.text_frame.add_paragraph().text = "Two"

    result = ingest_pptx(_deck(tmp_path, build), "blk")
    block = tmp_path / "blocks" / "blk"
    write_block(result, block, "deck.pptx")
    for rel in result.media:
        assert (block / rel).is_file()
    out = render_session(
        session_path=block / "slides.md",
        template_path=EXAMPLES / "template.pptx",
        layout_map_path=EXAMPLES / "layout-map.yaml",
        out_dir=tmp_path / "out",
    )
    with zipfile.ZipFile(out["pptx"]) as z:
        media_parts = [n for n in z.namelist() if n.startswith("ppt/media/")]
    assert len(media_parts) >= 2
