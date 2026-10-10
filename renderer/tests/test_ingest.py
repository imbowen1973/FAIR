"""Ingestion: an existing .pptx becomes a block of grammar + clean media.

The fixture deck is deliberately a legacy horror: free-positioned text
boxes (no placeholders), a manually-bulleted list, two columns, a
picture, a table, SmartArt-free but chart-free too — the common case.
The strongest assertion is the round trip: what the import emits must
parse and RENDER through the real pipeline without edits.
"""

from __future__ import annotations

import io
from pathlib import Path

import pytest
import yaml

from edufair_renderer.ingest import ingest_pptx, write_block
from edufair_renderer.parser import parse_session
from edufair_renderer.render import render_session

REPO = Path(__file__).resolve().parents[2]
EXAMPLES = REPO / "examples"


@pytest.fixture(scope="module")
def legacy_deck(tmp_path_factory) -> Path:
    from PIL import Image
    from pptx import Presentation
    from pptx.util import Emu, Inches, Pt

    prs = Presentation()
    blank = prs.slide_layouts[6]

    # 1: a hand-made title slide — one big text box, no placeholder.
    s = prs.slides.add_slide(blank)
    tb = s.shapes.add_textbox(Inches(1), Inches(0.6), Inches(8), Inches(1.2))
    run = tb.text_frame.paragraphs[0].add_run()
    run.text = "Cold chain handling"
    run.font.size = Pt(40)

    # 2: title + manually bulleted wall of text.
    s = prs.slides.add_slide(blank)
    t = s.shapes.add_textbox(Inches(0.8), Inches(0.4), Inches(8), Inches(0.9))
    r = t.text_frame.paragraphs[0].add_run()
    r.text = "Receiving goods"
    r.font.size = Pt(28)
    body = s.shapes.add_textbox(Inches(0.8), Inches(1.8), Inches(8), Inches(4))
    tf = body.text_frame
    first = tf.paragraphs[0].add_run()
    first.text = "• Check the logger"
    p2 = tf.add_paragraph()
    p2.level = 1
    p2.add_run().text = "- Thresholds first"
    p3 = tf.add_paragraph()
    run_bold = p3.add_run()
    run_bold.text = "Quarantine on breach"
    run_bold.font.bold = True
    s.notes_slide.notes_text_frame.text = "Walk the dock with the group."

    # 3: two free columns.
    s = prs.slides.add_slide(blank)
    t = s.shapes.add_textbox(Inches(0.8), Inches(0.4), Inches(8), Inches(0.9))
    t.text_frame.paragraphs[0].add_run().text = "Paper vs digital"
    left = s.shapes.add_textbox(Inches(0.6), Inches(1.8), Inches(3.8), Inches(3))
    left.text_frame.paragraphs[0].add_run().text = "Slow audits"
    right = s.shapes.add_textbox(Inches(5.2), Inches(1.8), Inches(3.8), Inches(3))
    right.text_frame.paragraphs[0].add_run().text = "Instant audit packs"

    # 4: a picture with a caption-sized note.
    img_path = tmp_path_factory.mktemp("img") / "big.png"
    Image.new("RGB", (2600, 1400), (30, 90, 140)).save(img_path)
    s = prs.slides.add_slide(blank)
    t = s.shapes.add_textbox(Inches(0.8), Inches(0.4), Inches(8), Inches(0.9))
    t.text_frame.paragraphs[0].add_run().text = "A week of readings"
    s.shapes.add_picture(str(img_path), Inches(1.5), Inches(1.6), width=Inches(6))
    cap = s.shapes.add_textbox(Inches(1.5), Inches(5.6), Inches(6), Inches(0.6))
    cap.text_frame.paragraphs[0].add_run().text = "Red line marks the threshold."

    # 5: a table.
    s = prs.slides.add_slide(blank)
    t = s.shapes.add_textbox(Inches(0.8), Inches(0.4), Inches(8), Inches(0.9))
    t.text_frame.paragraphs[0].add_run().text = "Alarm thresholds"
    frame = s.shapes.add_table(3, 2, Inches(1), Inches(1.8), Inches(7), Inches(2))
    cells = [["Product", "Max °C"], ["Vaccine", "8"], ["Insulin", "25"]]
    for ri, row in enumerate(cells):
        for ci, val in enumerate(row):
            frame.table.cell(ri, ci).text = val

    out = tmp_path_factory.mktemp("deck") / "legacy.pptx"
    prs.save(str(out))
    return out


@pytest.fixture(scope="module")
def result(legacy_deck):
    return ingest_pptx(legacy_deck, "05-imported")


def test_every_slide_mapped(result):
    assert [r["layout"] for r in result.report] == [
        "Section", "Full", "Split", "Picture", "Full",
    ]
    assert all(r["confidence"] in ("high", "medium", "low") for r in result.report)


def test_bullets_cleaned_and_emphasis_kept(result):
    slides = {s["id"]: s for s in _slides(result.slides_md)}
    items = slides["s-02"]["full"]["items"]
    assert items[0] == {"text": "Check the logger", "items": ["Thresholds first"]}
    assert items[1] == "**Quarantine on breach**"
    assert slides["s-02"]["notes"].startswith("Walk the dock")


def test_columns_become_split(result):
    slides = {s["id"]: s for s in _slides(result.slides_md)}
    assert slides["s-03"]["left"]["text"] == "Slow audits"
    assert slides["s-03"]["right"]["text"] == "Instant audit packs"


def test_picture_through_policy(result):
    slides = {s["id"]: s for s in _slides(result.slides_md)}
    src = slides["s-04"]["picture"]["src"]
    assert src in result.media
    from PIL import Image

    im = Image.open(io.BytesIO(result.media[src]))
    assert max(im.size) <= 2000  # resized at the door
    assert slides["s-04"]["caption"]["text"] == "Red line marks the threshold."


def test_table_extracted(result):
    slides = {s["id"]: s for s in _slides(result.slides_md)}
    table = slides["s-05"]["full"]
    assert table["type"] == "table"
    assert table["header"] == ["Product", "Max °C"]
    assert table["rows"] == [["Vaccine", "8"], ["Insulin", "25"]]


def test_round_trip_renders(result, tmp_path):
    """The import's output renders through the real pipeline, unedited."""
    block = tmp_path / "blocks" / "05-imported"
    write_block(result, block, "legacy.pptx")
    rendered = render_session(
        session_path=block / "slides.md",
        template_path=EXAMPLES / "template.pptx",
        layout_map_path=EXAMPLES / "layout-map.yaml",
        out_dir=tmp_path / "out",
    )
    assert rendered["pptx"].exists()
    assert (block / "report.md").read_text().count("|") > 10
    assert (block / "lessonplan.md").exists()


def test_table_type_renders_native_table(tmp_path):
    """The new table content type produces a real <a:tbl> in the deck."""
    import zipfile

    md = tmp_path / "s.md"
    md.write_text(
        "---\nsession: t1\ntitle: T\nversion: 0.1.0\n---\n\n"
        "--- slide\nid: s-01\nlayout: Full\ntitle: Thresholds\n"
        "full:\n  type: table\n  header: [Product, Max]\n"
        "  rows:\n  - [Vaccine, '8']\n  - [Insulin, '25']\n---\n",
        encoding="utf-8",
    )
    rendered = render_session(
        session_path=md,
        template_path=EXAMPLES / "template.pptx",
        layout_map_path=EXAMPLES / "layout-map.yaml",
        out_dir=tmp_path / "out",
    )
    with zipfile.ZipFile(rendered["pptx"]) as z:
        xml = z.read("ppt/slides/slide1.xml").decode()
    assert "<a:tbl>" in xml
    assert "Vaccine" in xml and "Insulin" in xml


def _slides(slides_md: str) -> list[dict]:
    session = parse_session_text(slides_md)
    return session


def parse_session_text(text: str) -> list[dict]:
    """Raw YAML slide dicts, without going through the full parser types."""
    blocks = []
    chunks = text.split("--- slide")
    for chunk in chunks[1:]:
        body = chunk.split("\n---", 1)[0]
        blocks.append(yaml.safe_load(body))
    return blocks
