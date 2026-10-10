"""SmartArt: import keeps the content tree, the render gives PowerPoint
real SmartArt back, and the generated data model is PowerPoint's own."""

from __future__ import annotations

import gzip
import io
import zipfile
from pathlib import Path

import pytest
from lxml import etree

from edufair_renderer import smartart as sa
from edufair_renderer.ingest import ingest_pptx, write_block
from edufair_renderer.parser import SessionParseError, SmartArtNode, parse_session
from edufair_renderer.render import render_session

REPO = Path(__file__).resolve().parents[2]
EXAMPLES = REPO / "examples"
FIXTURES = Path(__file__).parent / "fixtures" / "smartart"
LAYOUTS = sorted(p.name.split(".")[0] for p in FIXTURES.glob("*.definition.xml.gz"))
DGM = sa.DGM


def definition(name: str) -> sa.Definition:
    return sa.read_definition(gzip.decompress((FIXTURES / f"{name}.definition.xml.gz").read_bytes()))


def powerpoint_data(name: str) -> bytes:
    return gzip.decompress((FIXTURES / f"{name}.data.xml.gz").read_bytes())


def pres_tree(data: bytes) -> list[str]:
    """A data model's presentation tree, free of ids: what each shape is
    called, what it belongs to, what it shows and how it is styled."""
    root = etree.fromstring(data)
    A = sa.A
    pts = {p.get("modelId"): p for p in root.iter(f"{{{DGM}}}pt")}
    cxns = list(root.iter(f"{{{DGM}}}cxn"))
    owner = {}
    for c in cxns:
        if (c.get("type") or "parOf") == "parOf":
            owner[c.get("parTransId")] = c.get("destId")
            owner[c.get("sibTransId")] = c.get("destId")

    def text(pid):
        return "|".join(
            "".join(t.text or "" for t in p.iter(f"{{{A}}}t")) for p in pts[pid].iter(f"{{{A}}}p")
        ).strip("|")

    def label(pid):
        kind = pts[pid].get("type") or "node"
        if kind == "doc":
            return "DOC"
        if kind in ("parTrans", "sibTrans"):
            return f"{kind}({text(owner[pid])})"
        return f"{kind}:{text(pid)}"

    kids, has_parent, shows = {}, set(), {}
    for c in cxns:
        kind = c.get("type") or "parOf"
        if kind == "presParOf":
            kids.setdefault(c.get("srcId"), []).append((int(c.get("srcOrd") or 0), c.get("destId")))
            has_parent.add(c.get("destId"))
        elif kind == "presOf":
            shows.setdefault(c.get("destId"), []).append((int(c.get("destOrd") or 0), label(c.get("srcId"))))
    out = []

    def walk(pid, depth):
        ps = pts[pid].find(f"{{{DGM}}}prSet")
        out.append(
            "  " * depth
            + f"{ps.get('presName')} @{label(ps.get('presAssocID'))} {ps.get('presStyleLbl')} "
            f"{ps.get('presStyleIdx')}/{ps.get('presStyleCnt')} "
            f"[{';'.join(x for _, x in sorted(shows.get(pid, [])))}]"
        )
        for _, child in sorted(kids.get(pid, [])):
            walk(child, depth + 1)

    for pid, p in pts.items():
        if p.get("type") == "pres" and pid not in has_parent:
            walk(pid, 0)
    return out


# ---- the generator against PowerPoint ---------------------------------------


@pytest.mark.parametrize("name", LAYOUTS)
def test_presentation_points_match_powerpoint(name):
    """From the content alone, the walk of the layout definition rebuilds
    exactly the presentation tree PowerPoint wrote: every shape, what it
    shows, its style label and colour index."""
    want = pres_tree(powerpoint_data(name))
    got = pres_tree(sa.regenerate(powerpoint_data(name), definition(name)))
    assert want, "the fixture has no presentation points"
    assert got == want


def test_families_from_the_layouts_own_categories():
    families = {name: definition(name).family for name in LAYOUTS}
    assert families == {
        "cycle2": "cycle", "default": "list", "hierarchy4": "hierarchy", "list1": "list",
        "matrix1": "matrix", "orgChart1": "hierarchy", "process3": "process",
        # declares relationship and cycle; cycle has the leading priority
        "radial3": "cycle",
    }


def test_definition_bundle_round_trips():
    d = definition("process3")
    again = sa.read_definition(d.to_bytes())
    assert again.to_bytes() == d.to_bytes()
    assert again.layout_id == "urn:microsoft.com/office/officeart/2005/8/layout/process3"
    with pytest.raises(sa.SmartArtError):
        sa.read_definition(b"<notours/>")


def test_edited_content_gets_its_own_shapes():
    """Add a step: the generated model grows a shape for it, which is
    what lets PowerPoint lay out content it never saw."""
    d = definition("cycle2")
    nodes = [SmartArtNode(t) for t in "abcdef"]
    tree = pres_tree(sa.build_data_model(nodes, d, "seed"))
    assert sum(1 for line in tree if line.strip().startswith("node @node:")) == 6
    assert any("node @node:f node1 5/6" in line for line in tree)


def test_assistants_and_paragraphs_survive_the_model():
    nodes, _ = sa._parse_content(etree.fromstring(powerpoint_data("orgChart1")))
    manager = nodes[0]
    assert manager.text == "Manager\nSecond para"
    assert [c.text for c in manager.children] == ["Employee", "Employee2", "Assistant"]
    assert [c.assistant for c in manager.children] == [False, False, True]


# ---- the grammar ------------------------------------------------------------


def _session(tmp_path: Path, region: str) -> Path:
    md = tmp_path / "s.md"
    md.write_text(
        "---\nsession: t1\ntitle: T\nversion: 0.1.0\n---\n\n"
        "--- slide\nid: s-01\nlayout: Full\ntitle: Cold chain\n" + region + "---\n",
        encoding="utf-8",
    )
    return md


def _definition_file(tmp_path: Path, name: str = "process3") -> str:
    media = tmp_path / "media"
    media.mkdir(exist_ok=True)
    (media / f"smartart-{name}.xml").write_bytes(definition(name).to_bytes())
    return f"media/smartart-{name}.xml"


REGION = (
    "full:\n  type: smartart\n  layout: process3\n  family: process\n"
    "  definition: {src}\n  nodes:\n  - text: Receive\n    items: [Check logger]\n"
    "  - Verify\n  - Store at **2–8 °C**\n"
)


def test_parses_a_smartart_region(tmp_path):
    md = _session(tmp_path, REGION.format(src="media/x.xml"))
    region = parse_session(md).slides[0].regions[-1]
    assert region.type == "smartart" and region.src == "media/x.xml"
    assert region.family == "process" and region.layout == "process3"
    assert [n.text for n in region.nodes] == ["Receive", "Verify", "Store at **2–8 °C**"]
    assert region.nodes[0].children[0].text == "Check logger"


@pytest.mark.parametrize("bad, message", [
    ("full:\n  type: smartart\n  nodes: [a]\n", "requires 'definition'"),
    ("full:\n  type: smartart\n  definition: m.xml\n  nodes: []\n", "non-empty list"),
    ("full:\n  type: smartart\n  definition: m.xml\n  family: blob\n  nodes: [a]\n", "family"),
    ("full:\n  type: smartart\n  definition: m.xml\n  nodes:\n  - {text: a, colour: red}\n",
     "unknown key"),
])
def test_refuses_malformed_smartart(tmp_path, bad, message):
    with pytest.raises(SessionParseError, match=message):
        parse_session(_session(tmp_path, bad))


# ---- the render -------------------------------------------------------------


def _render(tmp_path, md, out="out"):
    return render_session(
        session_path=md,
        template_path=EXAMPLES / "template.pptx",
        layout_map_path=EXAMPLES / "layout-map.yaml",
        out_dir=tmp_path / out,
    )["pptx"]


def test_renders_real_smartart(tmp_path):
    md = _session(tmp_path, REGION.format(src=_definition_file(tmp_path)))
    pptx = _render(tmp_path, md)
    with zipfile.ZipFile(pptx) as z:
        names = z.namelist()
        slide = z.read("ppt/slides/slide1.xml").decode()
        rels = z.read("ppt/slides/_rels/slide1.xml.rels").decode()
        types = z.read("[Content_Types].xml").decode()
        data = z.read("ppt/diagrams/data1.xml")
    for part in ("data1", "layout1", "quickStyle1", "colors1"):
        assert f"ppt/diagrams/{part}.xml" in names
    assert not any("drawing" in n and "diagrams" in n for n in names)
    assert f'uri="{DGM}"' in slide
    for rel in ("diagramData", "diagramLayout", "diagramQuickStyle", "diagramColors"):
        assert f"relationships/{rel}" in rels
    assert "drawingml.diagramData+xml" in types
    # Content, emphasis and the shapes PowerPoint will draw.
    root = etree.fromstring(data)
    texts = ["".join(t.itertext()) for t in root.iter(f"{{{sa.A}}}t")]
    assert {"Receive", "Check logger", "Verify", "2–8 °C"} <= set(texts)
    bold = [r for r in root.iter(f"{{{sa.A}}}r") if r.find(f"{{{sa.A}}}rPr").get("b") == "1"]
    assert ["".join(r.itertext()) for r in bold] == ["2–8 °C"]
    tree = pres_tree(data)
    assert tree[0].startswith("linearFlow @DOC")
    assert sum("parTx @node:" in line for line in tree) == 3


def test_render_is_byte_deterministic(tmp_path):
    md = _session(tmp_path, REGION.format(src=_definition_file(tmp_path)))
    assert _render(tmp_path, md, "a").read_bytes() == _render(tmp_path, md, "b").read_bytes()


def test_missing_definition_is_named(tmp_path):
    from edufair_renderer.render import RenderError

    md = _session(tmp_path, REGION.format(src="media/gone.xml"))
    with pytest.raises(RenderError, match="media/gone.xml"):
        _render(tmp_path, md)


# ---- import, and the full circle ---------------------------------------------


def _deck_with_smartart(name: str, nodes, beside: str | None = None) -> bytes:
    """A deck holding one SmartArt, written by the same code the render
    uses -- the generated model is PowerPoint's own shape, so import
    meets exactly what PowerPoint would have saved."""
    from pptx import Presentation
    from pptx.util import Emu

    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[5])  # title only
    slide.shapes.title.text = "Cold chain"
    width = prs.slide_width
    left = Emu(width // 2) if beside else Emu(width // 10)
    sa.add_to_slide(slide, left, Emu(1600000), Emu(width * 4 // 10), Emu(4000000),
                    nodes, definition(name), "deck")
    if beside:
        box = slide.shapes.add_textbox(Emu(width // 20), Emu(1600000), Emu(width * 4 // 10), Emu(3000000))
        box.text_frame.text = beside
        box.text_frame.add_paragraph().text = "and a second line"
    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue()


def _region(result, key="full"):
    import yaml

    body = result.slides_md.split("--- slide\n", 1)[1].split("\n---", 1)[0]
    return yaml.safe_load(body)[key]


def test_import_keeps_smartart_whole():
    nodes = [SmartArtNode("Receive", [SmartArtNode("Check logger")]),
             SmartArtNode("Verify"), SmartArtNode("Store\nat 2–8 °C")]
    result = ingest_pptx(_deck_with_smartart("process3", nodes), "b", source_name="cc.pptx")
    region = _region(result)
    assert region["type"] == "smartart"
    assert region["layout"] == "process3" and region["family"] == "process"
    assert region["nodes"] == [{"text": "Receive", "items": ["Check logger"]}, "Verify",
                               "Store\nat 2–8 °C"]
    assert region["definition"] in result.media
    assert result.media[region["definition"]] == definition("process3").to_bytes()
    assert any("SmartArt kept as SmartArt: process3" in n for n in result.report[0]["notes"])
    # A node of two paragraphs is written as a block, not a quoted string.
    assert "  - |-\n    Store\n    at 2–8 °C" in result.slides_md


def test_import_puts_text_beside_the_diagram():
    nodes = [SmartArtNode("a"), SmartArtNode("b")]
    result = ingest_pptx(_deck_with_smartart("cycle2", nodes, beside="Why it loops"), "b")
    assert _region(result, "right")["type"] == "smartart"
    assert _region(result, "left")["type"] == "ul"


def test_import_render_import_is_the_identity(tmp_path):
    """Import a SmartArt, render the block, import the rendered deck: the
    same content and the same definition come back."""
    nodes = [SmartArtNode("Head", [SmartArtNode("One"), SmartArtNode("Two"),
                                   SmartArtNode("Aide", assistant=True)]),
             SmartArtNode("Peer")]
    first = ingest_pptx(_deck_with_smartart("orgChart1", nodes), "b", source_name="org.pptx")
    block = tmp_path / "blocks" / "b"
    write_block(first, block, "org.pptx")
    pptx = _render(tmp_path, block / "slides.md")
    again = ingest_pptx(pptx.read_bytes(), "b", source_name="org.pptx")
    assert _region(again) == _region(first)
    assert sorted(again.media) == sorted(first.media)
