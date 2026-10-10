"""Tables that stay in their frame.

PowerPoint grows a table row to fit its text and has no autofit for
tables, so a table rendered at the style's default size spills out of
its frame. These tests hold the renderer to the fix: the table's total
height fits the placeholder, the type size comes down as content goes
up, columns follow the source's proportions when it gives them, and an
imported deck's own proportions survive the round trip.
"""

from __future__ import annotations

import zipfile
from pathlib import Path

from lxml import etree

from edufair_renderer.render import render_session
from edufair_renderer.tablefit import EMU_PER_PT, MAX_PT, MIN_PT, fit_table

REPO = Path(__file__).resolve().parents[2]
EXAMPLES = REPO / "examples"
A_NS = "http://schemas.openxmlformats.org/drawingml/2006/main"
P_NS = "http://schemas.openxmlformats.org/presentationml/2006/main"

W = 8 * 914400  # an 8in x 4.5in frame
H = int(4.5 * 914400)

SENTENCE = "Quarantine the affected stock and record the peak temperature reached"


def test_small_table_stays_at_full_size():
    fit = fit_table(["Product", "Max °C"], [["Vaccine", "8"], ["Insulin", "25"]], W, H)
    assert fit.font_pt == MAX_PT
    assert fit.height <= H
    assert not fit.overflows


def test_dense_table_shrinks_but_fits():
    rows = [[f"Step {i}", SENTENCE, SENTENCE] for i in range(8)]
    fit = fit_table(["Step", "Action", "Rationale"], rows, W, H)
    assert MIN_PT <= fit.font_pt < MAX_PT
    assert fit.height <= H
    assert not fit.overflows


def test_more_content_never_gets_a_larger_font():
    sizes = []
    for n in (2, 6, 12):
        rows = [[f"Row {i}", SENTENCE] for i in range(n)]
        sizes.append(fit_table(["A", "B"], rows, W, H).font_pt)
    assert sizes == sorted(sizes, reverse=True)


def test_hopeless_table_reports_overflow_at_minimum():
    rows = [[SENTENCE * 3, SENTENCE * 3] for _ in range(30)]
    fit = fit_table(None, rows, W, H)
    assert fit.font_pt == MIN_PT
    assert fit.overflows


def test_columns_follow_given_widths_and_sum_exactly():
    fit = fit_table(None, [["a", "b", "c"]], W, H, widths=[50, 30, 20])
    assert sum(fit.column_widths) == W
    shares = [w / W for w in fit.column_widths]
    assert abs(shares[0] - 0.5) < 0.01 and abs(shares[2] - 0.2) < 0.01


def test_content_sizing_gives_the_wordy_column_more_room():
    rows = [["3", SENTENCE], ["4", SENTENCE]]
    fit = fit_table(["#", "What happens"], rows, W, H)
    assert fit.column_widths[1] > 3 * fit.column_widths[0]


def test_narrow_column_never_costs_the_table_its_size():
    """A short header in a narrow column must not drag the type down:
    columns get what their longest word needs before any sharing."""
    rows = [[str(i), "Confirm the logger reading against a second thermometer"]
            for i in range(1, 4)]
    fit = fit_table(["Step", "Action"], rows, W, H)
    assert fit.font_pt == MAX_PT


def test_fit_is_deterministic():
    rows = [[f"Row {i}", SENTENCE] for i in range(7)]
    assert fit_table(["A", "B"], rows, W, H) == fit_table(["A", "B"], rows, W, H)


def _render(tmp_path, region_yaml: str) -> etree._Element:
    md = tmp_path / "s.md"
    md.write_text(
        "---\nsession: t1\ntitle: T\nversion: 0.1.0\n---\n\n"
        "--- slide\nid: s-01\nlayout: Full\ntitle: A dense table\n"
        f"full:\n{region_yaml}---\n",
        encoding="utf-8",
    )
    out = render_session(
        session_path=md,
        template_path=EXAMPLES / "template.pptx",
        layout_map_path=EXAMPLES / "layout-map.yaml",
        out_dir=tmp_path / "out",
    )
    with zipfile.ZipFile(out["pptx"]) as z:
        return etree.fromstring(z.read("ppt/slides/slide1.xml"))


def test_rendered_dense_table_fits_its_placeholder(tmp_path):
    rows = "".join(
        f"  - ['Step {i}', '{SENTENCE}', '{SENTENCE}']\n" for i in range(8)
    )
    root = _render(
        tmp_path,
        "  type: table\n  header: [Step, Action, Rationale]\n  rows:\n" + rows,
    )

    frame = root.find(f".//{{{P_NS}}}graphicFrame")
    ext = frame.find(f"{{{P_NS}}}xfrm/{{{A_NS}}}ext")
    heights = [int(tr.get("h")) for tr in frame.iter(f"{{{A_NS}}}tr")]
    # Rows carry one shared one-line minimum (PowerPoint grows them to
    # their content); the frame carries the estimated total, which must
    # fit the Full layout's body.
    from pptx import Presentation

    layout = next(
        lo for lo in Presentation(str(EXAMPLES / "template.pptx")).slide_layouts
        if lo.name == "Full"
    )
    body = next(ph for ph in layout.placeholders if ph.placeholder_format.idx == 1)
    assert len(set(heights)) == 1
    assert int(ext.get("cy")) <= body.height

    # Every run carries the fitted size, below the style's default.
    sizes = {int(r.get("sz")) for r in frame.iter(f"{{{A_NS}}}rPr") if r.get("sz")}
    assert len(sizes) == 1
    (sz,) = sizes
    assert MIN_PT * 100 <= sz < MAX_PT * 100

    # Columns are not equal: the content decided.
    cols = [int(g.get("w")) for g in frame.iter(f"{{{A_NS}}}gridCol")]
    assert cols[0] < cols[1]
    assert sum(cols) == int(ext.get("cx"))


def test_rendered_widths_are_honoured(tmp_path):
    root = _render(
        tmp_path,
        "  type: table\n  widths: [70, 30]\n  rows:\n  - [a, b]\n  - [c, d]\n",
    )
    cols = [int(g.get("w")) for g in root.iter(f"{{{A_NS}}}gridCol")]
    assert abs(cols[0] / sum(cols) - 0.7) < 0.01


def test_widths_are_validated(tmp_path):
    import pytest

    from edufair_renderer.parser import SessionParseError

    with pytest.raises(SessionParseError, match="widths"):
        _render(tmp_path, "  type: table\n  widths: [1]\n  rows:\n  - [a, b]\n")


def test_import_keeps_the_decks_column_proportions(tmp_path):
    from pptx import Presentation
    from pptx.util import Inches

    from edufair_renderer.ingest import ingest_pptx

    prs = Presentation()
    s = prs.slides.add_slide(prs.slide_layouts[6])
    t = s.shapes.add_textbox(Inches(0.8), Inches(0.4), Inches(8), Inches(0.9))
    t.text_frame.paragraphs[0].add_run().text = "Thresholds"
    frame = s.shapes.add_table(3, 2, Inches(1), Inches(1.8), Inches(8), Inches(2))
    frame.table.columns[0].width = Inches(6)
    frame.table.columns[1].width = Inches(2)
    for r, row in enumerate([["Product", "Max"], ["Vaccine", "8"], ["Insulin", "25"]]):
        for c, val in enumerate(row):
            frame.table.cell(r, c).text = val
    deck = tmp_path / "t.pptx"
    prs.save(str(deck))

    result = ingest_pptx(deck, "x")
    assert "widths:\n  - 75\n  - 25" in result.slides_md

    # And what it emits renders, with those proportions.
    md = tmp_path / "x.md"
    md.write_text(result.slides_md, encoding="utf-8")
    out = render_session(
        session_path=md,
        template_path=EXAMPLES / "template.pptx",
        layout_map_path=EXAMPLES / "layout-map.yaml",
        out_dir=tmp_path / "out",
    )
    with zipfile.ZipFile(out["pptx"]) as z:
        root = etree.fromstring(z.read("ppt/slides/slide1.xml"))
    cols = [int(g.get("w")) for g in root.iter(f"{{{A_NS}}}gridCol")]
    assert abs(cols[0] / sum(cols) - 0.75) < 0.01
