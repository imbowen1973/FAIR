"""edufair-import: rescue the knowledge out of an existing PowerPoint deck.

Reads any .pptx — tidy placeholders or free-positioned chaos — and emits
a block folder in the library format: `slides.md` in the session
grammar, `media/` put through the image policy, a `block.yaml` and
`lessonplan.md` stub, and `report.md` saying per slide how confident the
mapping is and what needs a human.

The job is to rescue the CONTENT — text, structure, images, notes — not
to replicate the old slide's appearance. The free-positioned arrangement
is what made legacy decks unreusable; shedding it is the conversion.
Extraction reads every shape regardless of placeholder membership;
classification works from geometry and type, because most legacy decks
never used placeholders at all.

    edufair-import old-course.pptx --out blocks/05-imported

The source deck is read and discarded; only markdown and policy-clean
images come out. Runs identically on the CLI and in the workbench's
browser runtime (Pyodide), so one extractor serves bulk migration and
drop-a-file-on-the-door alike.
"""

from __future__ import annotations

import hashlib
import io
import re
import sys
from dataclasses import dataclass, field
from pathlib import Path

import yaml

# Geometry heuristics, as fractions of the slide.
TITLE_BAND = 0.30  # a title lives in the top of the slide
SIDE_GAP = 0.22  # two bodies this far apart (centre to centre) are columns
CAPTION_CHARS = 220
SECTION_WORDS = 7

MAX_EDGE = 2000
JPEG_QUALITY = 82


@dataclass
class _Text:
    paragraphs: list[dict]  # {text, level}
    left: int
    top: int
    width: int
    height: int
    is_title_ph: bool = False
    font_size: int = 0  # EMU-agnostic: pt * 100, 0 = unknown

    @property
    def chars(self) -> int:
        return sum(len(p["text"]) for p in self.paragraphs)


@dataclass
class _Pic:
    blob: bytes
    ext: str
    left: int
    top: int
    width: int = 0
    height: int = 0
    # Sat in a picture/content placeholder: the strongest sign it is the
    # slide's picture rather than decoration.
    in_placeholder: bool = False
    # PowerPoint crops as fractions of the original (left, top, right,
    # bottom). The blob is uncropped; the crop is applied on import so the
    # media matches what the slide showed.
    crop: tuple[float, float, float, float] = (0.0, 0.0, 0.0, 0.0)
    # The slide's own background picture: saved, never placed -- the
    # template owns backgrounds.
    background: bool = False

    @property
    def area(self) -> int:
        return max(0, self.width) * max(0, self.height)


@dataclass
class _Facts:
    number: int
    texts: list[_Text] = field(default_factory=list)
    pics: list[_Pic] = field(default_factory=list)
    tables: list[dict] = field(default_factory=list)  # {header, rows}
    notes: str | None = None
    oddities: list[str] = field(default_factory=list)


@dataclass
class IngestResult:
    title: str
    slides_md: str
    media: dict[str, bytes]  # path relative to the block dir -> bytes
    report: list[dict]  # {slide, id, layout, confidence, notes}


# A typed bullet: a glyph FOLLOWED BY WHITESPACE. The space is what
# makes it a bullet -- without it, "oxygen" lost its "o" and "-5"
# became "5", which in a cold-chain table turns -20 °C into 20 °C.
_BULLET_GLYPHS = re.compile(r"^\s*(?:[•‣◦·⁃\-–—o]\s+)+")
_DIAGRAM_URI = "http://schemas.openxmlformats.org/drawingml/2006/diagram"
_A_T = "{http://schemas.openxmlformats.org/drawingml/2006/main}t"


def _clean(text: str) -> str:
    return _BULLET_GLYPHS.sub("", text).strip()


def _run_text(run) -> str:
    text = run.text or ""
    if not text.strip():
        return text
    if run.font.bold:
        return f"**{text}**"
    return text


def _paragraphs(text_frame) -> list[dict]:
    out = []
    for para in text_frame.paragraphs:
        text = _clean("".join(_run_text(r) for r in para.runs))
        if text:
            out.append({"text": text, "level": para.level or 0})
    return out


_R_EMBED = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}embed"
_A_BLIP = "{http://schemas.openxmlformats.org/drawingml/2006/main}blip"


def _blip_image(part, blip):
    """The image a <a:blip> points at, or None."""
    rid = blip.get(_R_EMBED) if blip is not None else None
    if not rid:
        return None
    try:
        image_part = part.related_part(rid)
        return image_part.blob, image_part.partname.ext
    except Exception:
        return None


def _picture_fill(shape, left: int, top: int) -> "_Pic | None":
    """A shape whose fill is a picture (<p:spPr><a:blipFill>)."""
    sppr = getattr(shape._element, "spPr", None)
    if sppr is None:
        return None
    blip = sppr.find(f".//{_A_BLIP}")
    found = _blip_image(shape.part, blip)
    if found is None:
        return None
    blob, ext = found
    return _Pic(blob=blob, ext=ext, left=left, top=top,
                width=shape.width or 0, height=shape.height or 0)


def _background_picture(slide, size) -> "_Pic | None":
    """The slide's own background picture (<p:bg>), if it has one.

    Only the slide's own: a background inherited from the layout or
    master is the old template's, and there is nothing to rescue."""
    bg = slide._element.find(
        "{http://schemas.openxmlformats.org/presentationml/2006/main}cSld/"
        "{http://schemas.openxmlformats.org/presentationml/2006/main}bg"
    )
    if bg is None:
        return None
    found = _blip_image(slide.part, bg.find(f".//{_A_BLIP}"))
    if found is None:
        return None
    blob, ext = found
    return _Pic(blob=blob, ext=ext, left=0, top=0, width=size[0], height=size[1],
                background=True)


def _table_grid(table) -> list[list[str]]:
    """A table's cells as the grammar writes them, merges as markers.

    A merge's text lives in its top-left cell; every cell it covers
    becomes "<" (same row as the top-left) or "^" (a row below), which
    the parser turns back into the same merge. Cells are not
    bullet-stripped -- they are not lists, and a leading "-" in a table
    is usually a sign -- and a cell that literally reads "<" or "^" is
    escaped so it is not taken for a marker.
    """
    n_rows, n_cols = len(table.rows), len(table.columns)
    grid: list[list[str | None]] = [[None] * n_cols for _ in range(n_rows)]
    for r in range(n_rows):
        for c in range(n_cols):
            if grid[r][c] is not None:
                continue  # already claimed by a merge above or to the left
            cell = table.cell(r, c)
            text = cell.text.strip()
            grid[r][c] = {"<": "\\<", "^": "\\^"}.get(text, text)
            if cell.is_merge_origin:
                for i in range(r, min(n_rows, r + cell.span_height)):
                    for j in range(c, min(n_cols, c + cell.span_width)):
                        if (i, j) != (r, c):
                            grid[i][j] = "<" if i == r else "^"
    return grid


def _font_size(text_frame) -> int:
    best = 0
    for para in text_frame.paragraphs:
        for run in para.runs:
            if run.font.size is not None:
                best = max(best, int(run.font.size.pt * 100))
    return best


def _harvest(shape, facts: _Facts, offset=(0, 0)) -> None:
    from pptx.enum.shapes import MSO_SHAPE_TYPE, PP_PLACEHOLDER

    kind = shape.shape_type
    left = (shape.left or 0) + offset[0]
    top = (shape.top or 0) + offset[1]

    if kind == MSO_SHAPE_TYPE.GROUP:
        # Children keep their own coordinate space; the group's position
        # is a fair approximation for classification.
        for child in shape.shapes:
            _harvest(child, facts, (left, top))
        return

    # Any picture: free-floating (Picture) or sitting in a placeholder
    # (PlaceholderPicture). The second is a different python-pptx type
    # with shape_type PLACEHOLDER, and looking only for PICTURE dropped
    # every placeholder image without a word.
    from pptx.shapes.picture import _BasePicture

    if isinstance(shape, _BasePicture):
        try:
            image = shape.image
            facts.pics.append(
                _Pic(
                    blob=image.blob, ext=image.ext, left=left, top=top,
                    width=shape.width or 0, height=shape.height or 0,
                    in_placeholder=shape.is_placeholder,
                    crop=(
                        shape.crop_left, shape.crop_top,
                        shape.crop_right, shape.crop_bottom,
                    ),
                )
            )
        except Exception:
            facts.oddities.append(
                "a picture could not be read (linked, not embedded?) and was dropped"
            )
        return

    # A picture used as a shape's FILL (a rectangle filled with a photo).
    # Its text, if any, is harvested below as ordinary text; the overlay
    # check then sees it sits on the picture.
    fill_pic = _picture_fill(shape, left, top)
    if fill_pic is not None:
        facts.pics.append(fill_pic)

    if getattr(shape, "has_table", False) and shape.has_table:
        rows = _table_grid(shape.table)
        header = rows[0] if len(rows) > 1 else None
        body = rows[1:] if len(rows) > 1 else rows
        # The deck's own column proportions, kept as whole percentages:
        # they say what the columns hold, which is content. The font size
        # is deliberately not kept -- it is usually inherited rather than
        # set, and a size that suited the old frame is wrong for the new
        # one. The renderer fits the type to the frame instead.
        raw_widths = [int(col.width or 0) for col in shape.table.columns]
        total = sum(raw_widths)
        widths = None
        if total > 0 and all(raw_widths) and len(set(raw_widths)) > 1:
            widths = [max(1, round(100 * w / total)) for w in raw_widths]
        facts.tables.append({"header": header, "rows": body, "widths": widths})
        return

    if getattr(shape, "has_chart", False) and shape.has_chart:
        facts.oddities.append(
            "a live chart was found; its data was not extracted — rebuild as a table"
        )
        return

    if kind == MSO_SHAPE_TYPE.PLACEHOLDER or shape.has_text_frame:
        element_xml = shape._element.xml if hasattr(shape, "_element") else ""
        if _DIAGRAM_URI in element_xml:
            texts = [
                _clean(t.text or "") for t in shape._element.iter(_A_T)
            ]
            texts = [t for t in texts if t]
            if texts:
                facts.texts.append(
                    _Text(
                        paragraphs=[{"text": t, "level": 0} for t in texts],
                        left=left, top=top,
                        width=shape.width or 0, height=shape.height or 0,
                    )
                )
            facts.oddities.append(
                "SmartArt flattened to a list — redraw as Mermaid if the shape mattered"
            )
            return
        if not shape.has_text_frame:
            return
        paragraphs = _paragraphs(shape.text_frame)
        if not paragraphs:
            return
        is_title = False
        if shape.is_placeholder:
            ph = shape.placeholder_format.type
            is_title = ph in (PP_PLACEHOLDER.TITLE, PP_PLACEHOLDER.CENTER_TITLE)
        facts.texts.append(
            _Text(
                paragraphs=paragraphs,
                left=left, top=top,
                width=shape.width or 0, height=shape.height or 0,
                is_title_ph=is_title,
                font_size=_font_size(shape.text_frame),
            )
        )


def extract(source) -> tuple[list[_Facts], tuple[int, int], str | None]:
    """Everything worth keeping, slide by slide, placeholders or not."""
    from pptx import Presentation

    prs = Presentation(io.BytesIO(source) if isinstance(source, bytes) else str(source))
    deck_title = None
    slides = []
    for n, slide in enumerate(prs.slides, start=1):
        facts = _Facts(number=n)
        for shape in slide.shapes:
            _harvest(shape, facts)
        background = _background_picture(slide, (prs.slide_width, prs.slide_height))
        if background is not None:
            facts.pics.append(background)
        if slide.has_notes_slide:
            notes = slide.notes_slide.notes_text_frame.text.strip()
            facts.notes = notes or None
        slides.append(facts)
        if deck_title is None:
            for t in facts.texts:
                if t.is_title_ph and t.paragraphs:
                    deck_title = t.paragraphs[0]["text"].replace("**", "")
                    break
    return slides, (prs.slide_width, prs.slide_height), deck_title


def _pick_title(facts: _Facts, slide_h: int) -> _Text | None:
    for t in facts.texts:
        if t.is_title_ph:
            return t
    in_band = [t for t in facts.texts if t.top < slide_h * TITLE_BAND]
    if not in_band:
        return None
    # Largest declared type wins; ties go to the topmost.
    in_band.sort(key=lambda t: (-t.font_size, t.top))
    best = in_band[0]
    if len(facts.texts) > 1 and best.chars > 120:
        return None  # a wall of text is content, not a title
    return best


def _items(paragraphs: list[dict]):
    """Nested grammar items from levelled paragraphs."""
    root: list = []
    stack = [(0, root)]
    for para in paragraphs:
        level = max(0, min(4, para["level"]))
        while stack and stack[-1][0] > level:
            stack.pop()
        while stack[-1][0] < level:
            siblings = stack[-1][1]
            if not siblings:
                break  # a child with no parent: flatten up
            parent = siblings[-1]
            if isinstance(parent, str):
                siblings[-1] = parent = {"text": parent, "items": []}
            elif "items" not in parent:
                parent["items"] = []
            stack.append((stack[-1][0] + 1, parent["items"]))
        stack[-1][1].append(para["text"])
    return root


def _text_region(body: _Text):
    if len(body.paragraphs) == 1 and body.chars <= CAPTION_CHARS:
        return {"type": "p", "text": body.paragraphs[0]["text"]}
    return {"type": "ul", "items": _items(body.paragraphs)}


def _columns(a: _Text, b: _Text, slide_w: int) -> bool:
    ca = a.left + (a.width or 0) / 2
    cb = b.left + (b.width or 0) / 2
    return abs(ca - cb) > slide_w * SIDE_GAP


def classify(facts: _Facts, size, media_ref) -> tuple[dict, str, list[str]]:
    """One slide's facts -> (slide dict, confidence, notes)."""
    slide_w, slide_h = size
    notes = list(facts.oddities)
    title = _pick_title(facts, slide_h)
    bodies = [t for t in facts.texts if t is not title]
    bodies.sort(key=lambda t: (t.left, t.top))
    out: dict = {}
    if title:
        out["title"] = title.paragraphs[0]["text"].replace("**", "")
        for extra in title.paragraphs[1:]:
            bodies.append(_Text(paragraphs=[extra], left=title.left, top=title.top,
                                width=title.width, height=title.height))

    def done(layout, confidence):
        out_full = {"layout": layout, **out}
        return out_full, confidence, notes

    if facts.tables:
        table = facts.tables[0]
        region = {"type": "table", "rows": table["rows"]}
        if table["header"]:
            region["header"] = table["header"]
        if table.get("widths"):
            region["widths"] = table["widths"]
        out["full"] = region
        if len(facts.tables) > 1:
            notes.append(f"{len(facts.tables) - 1} further table(s) dropped — split into their own slides")
        if bodies or facts.pics:
            notes.append("table kept; other content on this slide set aside")
        return done("Full", "medium")

    if not bodies and not facts.pics:
        if not title:
            notes.append("nothing extractable — slide skipped content")
            return done("TitleOnly", "low")
        words = len(out.get("title", "").split())
        return done("Section" if words <= SECTION_WORDS else "TitleOnly", "high")

    if not facts.pics:
        if len(bodies) == 1:
            out["full"] = _text_region(bodies[0])
            return done("Full", "high" if bodies[0].is_title_ph or title else "medium")
        if len(bodies) == 2 and _columns(bodies[0], bodies[1], slide_w):
            out["left"] = _text_region(bodies[0])
            out["right"] = _text_region(bodies[1])
            return done("Split", "medium")
        merged = [p for b in bodies for p in b.paragraphs]
        out["full"] = {"type": "ul", "items": _items(merged)}
        notes.append(f"{len(bodies)} text areas merged — free-form layout approximated")
        return done("Full", "low")

    # The picture that becomes the slide's: one that sat in a placeholder
    # if there is one, else the largest. Backgrounds never compete.
    placed = [pic for pic in facts.pics if not pic.background]
    for pic in facts.pics:
        if pic.background:
            kept = media_ref(pic, facts.number)
            if kept:
                notes.append(
                    f"slide background picture saved as {kept}, not placed — "
                    "the template now provides the background"
                )
    if not placed:
        # Only a background: the slide is its text.
        facts_without = _Facts(number=facts.number, texts=facts.texts,
                               tables=facts.tables, notes=facts.notes)
        slide_dict, confidence, more = classify(facts_without, size, media_ref)
        return slide_dict, confidence, notes + more

    main = max(placed, key=lambda pic: (pic.in_placeholder, pic.area))
    pic_ref = media_ref(main, facts.number)
    for pic in placed:
        if pic is not main:
            kept = media_ref(pic, facts.number)
            if kept:
                notes.append(f"extra picture saved as {kept} — place it on its own slide")

    if pic_ref is None:
        notes.append("picture in an unsupported format was dropped")
        if bodies:
            out["full"] = _text_region(bodies[0])
            return done("Full", "low")
        return done("TitleOnly", "low")

    # Text sitting ON the picture can no longer sit on it: the template
    # owns positions. It is kept -- as the caption, or with the body --
    # and the report quotes it so a reviewer can see what was overlaid.
    if title is not None and _overlap(title, main) >= OVERLAY_SHARE:
        notes.append(f"title \"{_snip(out.get('title', ''))}\" was set over the picture")
    overlaid = [b for b in bodies if _overlap(b, main) >= OVERLAY_SHARE]
    beside = [b for b in bodies if b not in overlaid]
    for b in overlaid:
        notes.append(f"text \"{_snip(_plain(b))}\" was overlaid on the picture")

    picture = {"type": "image", "src": pic_ref}

    if not beside:
        out["picture"] = picture
        if overlaid:
            out["caption"] = _merged_region(overlaid)
            notes.append("overlaid text kept as the caption")
        return done("Picture", "medium")

    # A single short line under (or over) the picture is its caption.
    # A list or several lines beside it is body text, and goes to the side.
    if (
        len(beside) == 1 and not overlaid
        and len(beside[0].paragraphs) == 1 and beside[0].chars <= CAPTION_CHARS
        and _horizontal_share(beside[0], main) >= 0.5
    ):
        out["picture"] = picture
        out["caption"] = _text_region(beside[0])
        return done("Picture", "medium")

    # Picture and text side by side, on the side the picture was.
    text_region = _merged_region(beside + overlaid)
    if len(beside) > 1:
        notes.append(f"{len(beside)} text areas merged beside the picture")
    if overlaid:
        notes.append("overlaid text moved beside the picture")
    picture_left = _centre_x(main) < sum(_centre_x(b) for b in beside) / len(beside)
    out["left"], out["right"] = (
        (picture, text_region) if picture_left else (text_region, picture)
    )
    clean_split = all(_horizontal_share(b, main) < 0.2 for b in beside)
    return done("Split", "medium" if clean_split else "low")


# Text overlapping the picture by at least this share of its own area
# was set ON the picture.
OVERLAY_SHARE = 0.5


def _box(shape) -> tuple[int, int, int, int]:
    return shape.left, shape.top, shape.left + shape.width, shape.top + shape.height


def _overlap(text: "_Text", pic: "_Pic") -> float:
    """Share of the text's area that lies on the picture."""
    tl, tt, tr, tb = _box(text)
    pl, pt, pr, pb = _box(pic)
    w = max(0, min(tr, pr) - max(tl, pl))
    h = max(0, min(tb, pb) - max(tt, pt))
    area = max(1, text.width * text.height)
    return (w * h) / area


def _horizontal_share(text: "_Text", pic: "_Pic") -> float:
    """Share of the text's width that sits within the picture's columns:
    high for a caption under a picture, low for text beside it."""
    tl, _, tr, _ = _box(text)
    pl, _, pr, _ = _box(pic)
    return max(0, min(tr, pr) - max(tl, pl)) / max(1, text.width)


def _centre_x(shape) -> float:
    return shape.left + shape.width / 2


def _plain(text: "_Text") -> str:
    return " ".join(p["text"].replace("**", "") for p in text.paragraphs)


def _snip(s: str, n: int = 50) -> str:
    return s if len(s) <= n else s[: n - 1] + "…"


def _merged_region(texts: list["_Text"]):
    """Several text areas as one region, in reading order."""
    texts = sorted(texts, key=lambda t: (t.top, t.left))
    paragraphs = [p for t in texts for p in t.paragraphs]
    if len(paragraphs) == 1 and len(paragraphs[0]["text"]) <= CAPTION_CHARS:
        return {"type": "p", "text": paragraphs[0]["text"]}
    return {"type": "ul", "items": _items(paragraphs)}


def _prepare_image(
    blob: bytes, ext: str, crop=(0.0, 0.0, 0.0, 0.0), cap: int | None = None
) -> tuple[bytes, str] | None:
    """The asset policy at the door: crop as shown, resize, re-encode,
    strip metadata, and fit under the repo's byte cap.

    Resizing alone is not enough: a 2000px photograph or a PNG with
    transparency can still exceed the cap, and the library's own build
    then refuses the commit. So quality comes down first (JPEG), then
    the image shrinks, until it fits."""
    from PIL import Image

    from .assets import MAX_IMAGE_BYTES

    cap = cap or MAX_IMAGE_BYTES
    try:
        im = Image.open(io.BytesIO(blob))
        im.load()
    except Exception:
        return None

    left, top, right, bottom = (max(0.0, float(c or 0)) for c in crop)
    if left or top or right or bottom:
        w, h = im.size
        box = (
            round(w * left), round(h * top),
            round(w * (1 - right)), round(h * (1 - bottom)),
        )
        if box[2] > box[0] and box[3] > box[1]:
            im = im.crop(box)

    edge = max(im.size)
    if edge > MAX_EDGE:
        scale = MAX_EDGE / edge
        im = im.resize((round(im.width * scale), round(im.height * scale)), Image.LANCZOS)
    has_alpha = im.mode in ("RGBA", "LA") or (im.mode == "P" and "transparency" in im.info)
    im = im.convert("RGBA" if has_alpha else "RGB")

    def encode(img, quality):
        buf = io.BytesIO()
        if has_alpha:
            img.save(buf, format="PNG", optimize=True)
        else:
            img.save(buf, format="JPEG", quality=quality, optimize=True, progressive=True)
        return buf.getvalue()

    for quality in (JPEG_QUALITY, 72, 62) if not has_alpha else (None,):
        data = encode(im, quality)
        if len(data) <= cap:
            return data, ("png" if has_alpha else "jpg")
    while min(im.size) > 64:
        im = im.resize((round(im.width * 0.85), round(im.height * 0.85)), Image.LANCZOS)
        data = encode(im, 62)
        if len(data) <= cap:
            break
    return data, ("png" if has_alpha else "jpg")


# An image on at least this many slides, and on at least this share of
# them, is the old deck's branding (a logo in the corner), not content.
BRANDING_MIN_SLIDES = 3
BRANDING_SHARE = 0.5


def _slug(text: str, limit: int = 24) -> str:
    slug = re.sub(r"[^a-z0-9]+", "-", (text or "").lower()).strip("-")
    return slug[:limit].strip("-") or "deck"


def ingest_pptx(
    source, block_id: str, title: str | None = None, source_name: str | None = None
) -> IngestResult:
    slides, size, deck_title = extract(source)
    media: dict[str, bytes] = {}
    by_key: dict[tuple, str] = {}

    # Media names say which deck and slide a picture came from, and end in
    # a content hash so they cannot collide. Numbering by slide alone
    # meant a second deck imported into the same block overwrote the
    # first deck's pictures, and its slides then showed the wrong image.
    if source_name is None and isinstance(source, (str, Path)):
        source_name = Path(source).name
    prefix = _slug(Path(source_name).stem if source_name else block_id)

    def media_ref(pic: _Pic, slide_no: int) -> str | None:
        key = (hashlib.sha1(pic.blob).hexdigest(), pic.crop)
        if key in by_key:
            return by_key[key]
        prepared = _prepare_image(pic.blob, pic.ext, crop=pic.crop)
        if prepared is None:
            return None
        data, ext = prepared
        name = f"media/{prefix}-s{slide_no:02d}-{hashlib.sha1(data).hexdigest()[:8]}.{ext}"
        media[name] = data
        by_key[key] = name
        return name

    # The old deck's branding -- the same picture on most slides -- is
    # the old template's job, which the new template now does. Dropped,
    # with one note, rather than placed on every slide as content.
    seen_on: dict[str, set[int]] = {}
    for facts in slides:
        for pic in facts.pics:
            seen_on.setdefault(hashlib.sha1(pic.blob).hexdigest(), set()).add(facts.number)
    branding = {
        digest for digest, numbers in seen_on.items()
        if len(numbers) >= BRANDING_MIN_SLIDES
        and len(numbers) >= BRANDING_SHARE * len(slides)
    }
    branding_noted = False
    if branding:
        for facts in slides:
            facts.pics = [
                pic for pic in facts.pics
                if hashlib.sha1(pic.blob).hexdigest() not in branding
            ]

    blocks: list[str] = []
    report: list[dict] = []
    first_title: str | None = None
    for facts in slides:
        slide_dict, confidence, notes = classify(facts, size, media_ref)
        if branding and not branding_noted:
            notes.insert(0, f"{len(branding)} picture(s) repeated on most slides "
                            "(a logo?) dropped as the old deck's branding")
            branding_noted = True
        if first_title is None and slide_dict.get("title"):
            first_title = slide_dict["title"]
        sid = f"s-{facts.number:02d}"
        ordered = {"id": sid, "layout": slide_dict.pop("layout")}
        ordered.update(slide_dict)
        if facts.notes:
            ordered["notes"] = facts.notes
        body = yaml.safe_dump(
            ordered, sort_keys=False, allow_unicode=True, width=76
        ).rstrip()
        blocks.append(f"--- slide\n{body}\n---")
        report.append(
            {"slide": facts.number, "id": sid, "layout": ordered["layout"],
             "confidence": confidence, "notes": notes}
        )

    final_title = title or deck_title or first_title or block_id
    front = yaml.safe_dump(
        {"session": block_id, "title": final_title, "version": "0.1.0"},
        sort_keys=False, allow_unicode=True,
    ).rstrip()
    slides_md = f"---\n{front}\n---\n\n" + "\n\n".join(blocks) + "\n"
    return IngestResult(title=final_title, slides_md=slides_md, media=media, report=report)


def write_block(result: IngestResult, out_dir: Path, source_name: str) -> None:
    out_dir.mkdir(parents=True, exist_ok=True)
    (out_dir / "slides.md").write_text(result.slides_md, encoding="utf-8")
    (out_dir / "block.yaml").write_text(
        yaml.safe_dump({"title": result.title}, sort_keys=False, allow_unicode=True),
        encoding="utf-8",
    )
    lesson = out_dir / "lessonplan.md"
    if not lesson.exists():
        lesson.write_text(
            f"# {result.title}\n\n*Imported from {source_name}. "
            "Write the lesson plan: the deck alone is not the session.*\n",
            encoding="utf-8",
        )
    for rel, data in result.media.items():
        target = out_dir / rel
        target.parent.mkdir(parents=True, exist_ok=True)
        target.write_bytes(data)

    low = sum(1 for r in result.report if r["confidence"] == "low")
    lines = [
        f"# Import report: {source_name}",
        "",
        f"{len(result.report)} slides; {low} need a human "
        "(low confidence). Review before submitting.",
        "",
        "| Slide | Id | Layout | Confidence | Notes |",
        "|---|---|---|---|---|",
    ]
    for r in result.report:
        lines.append(
            f"| {r['slide']} | {r['id']} | {r['layout']} | {r['confidence']} | "
            f"{'; '.join(r['notes']) or '—'} |"
        )
    (out_dir / "report.md").write_text("\n".join(lines) + "\n", encoding="utf-8")


def main(argv: list[str] | None = None) -> int:
    import argparse

    ap = argparse.ArgumentParser(
        prog="edufair-import",
        description="Rescue an existing .pptx into a block of markdown + policy-clean media",
    )
    ap.add_argument("pptx", type=Path)
    ap.add_argument("--out", type=Path, default=None,
                    help="block folder to write (default: blocks/<deck name>)")
    ap.add_argument("--id", dest="block_id", default=None,
                    help="block id (default: the deck's file name, slugified)")
    ap.add_argument("--title", default=None)
    args = ap.parse_args(argv)

    if not args.pptx.exists():
        print(f"error: {args.pptx} not found", file=sys.stderr)
        return 1
    slug = args.block_id or re.sub(r"[^a-z0-9]+", "-", args.pptx.stem.lower()).strip("-")
    out_dir = args.out or Path("blocks") / slug

    result = ingest_pptx(args.pptx, slug, title=args.title, source_name=args.pptx.name)
    write_block(result, out_dir, args.pptx.name)

    low = sum(1 for r in result.report if r["confidence"] == "low")
    print(f"{args.pptx.name}: {len(result.report)} slides -> {out_dir}")
    print(f"  {len(result.media)} image(s) through the policy; "
          f"{low} slide(s) flagged for review — see {out_dir / 'report.md'}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
